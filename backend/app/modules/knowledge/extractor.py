"""AI-powered content extraction for knowledge assets."""

import io
import structlog

logger = structlog.get_logger(__name__)


class ContentExtractor:
    """Extracts structured content from various file formats."""

    async def extract_from_pdf(self, file_bytes: bytes) -> dict:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return {"raw_text": text, "page_count": len(reader.pages)}

    async def extract_from_docx(self, file_bytes: bytes) -> dict:
        from docx import Document

        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return {"raw_text": "\n".join(paragraphs), "paragraph_count": len(paragraphs)}

    async def extract_from_pptx(self, file_bytes: bytes) -> dict:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(file_bytes))
        slides_text = []
        for slide in prs.slides:
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            slides_text.append(" ".join(slide_text))
        return {"raw_text": "\n".join(slides_text), "slide_count": len(prs.slides)}

    async def extract_text(self, file_bytes: bytes, content_type: str) -> dict:
        """Dispatch to appropriate extractor based on content_type."""
        extractors = {
            "pdf": self.extract_from_pdf,
            "docx": self.extract_from_docx,
            "pptx": self.extract_from_pptx,
        }
        extractor_fn = extractors.get(content_type)
        if extractor_fn:
            return await extractor_fn(file_bytes)
        # Plain text fallback
        return {"raw_text": file_bytes.decode("utf-8", errors="replace")}

    async def analyze_with_ai(self, text: str, orchestration_service) -> dict:
        """Use AI to extract structured topics, concepts, outcomes, and keywords."""
        prompt = (
            "Analyze the following text and extract:\n"
            "1. Main topics (list of strings)\n"
            "2. Key concepts (list of strings)\n"
            "3. Learning outcomes (list of strings)\n"
            "4. Keywords (list of strings)\n\n"
            "Respond ONLY in valid JSON with keys: topics, concepts, outcomes, keywords.\n\n"
            f"Text:\n{text[:4000]}"
        )
        messages = [{"role": "user", "content": prompt}]
        try:
            response = await orchestration_service.complete(messages)
            import json
            return json.loads(response)
        except Exception as exc:
            logger.warning("AI extraction failed", error=str(exc))
            return {"topics": [], "concepts": [], "outcomes": [], "keywords": []}
